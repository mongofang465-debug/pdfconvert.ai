import streamlit as st
import fitz  # PyMuPDF
from docx import Document
import camelot
import io
from zipfile import ZipFile
from openpyxl import Workbook

st.title("📄 PDF 转 Word / Excel / 图片 工具")
st.write("上传 PDF 文件，选择转换格式，一键下载处理后的文件")

# 上传文件
uploaded_files = st.file_uploader(
    "上传 PDF 文件（可多选）", type="pdf", accept_multiple_files=True
)

convert_option = st.selectbox(
    "选择转换类型",
    ("Word (.docx)", "Excel (.xlsx)", "图片 (.png 每页)")
)

# PDF 转 Word（PyMuPDF）
def pdf_to_word(pdf_file):
    doc = Document()
    pdf_bytes = pdf_file.read()
    pdf = fitz.open(stream=pdf_bytes, filetype="pdf")
    for page in pdf:
        text = page.get_text()
        if text:
            doc.add_paragraph(text)
    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return output

# PDF 转 Excel（文本+表格自动判断）
def pdf_to_excel(pdf_file):
    pdf_bytes = pdf_file.read()
    # 先尝试 Camelot 提取表格
    try:
        tables = camelot.read_pdf(io.BytesIO(pdf_bytes), pages='all')
    except:
        tables = []

    wb = Workbook()
    ws = wb.active

    if tables and len(tables) > 0:
        # 有表格时，使用表格数据
        for table in tables:
            for row in table.df.values.tolist():
                ws.append(row)
    else:
        # 没表格时，按文本每行写入
        pdf = fitz.open(stream=pdf_bytes, filetype="pdf")
        row_num = 1
        for page in pdf:
            text = page.get_text()
            lines = text.split("\n")
            for line in lines:
                if line.strip():
                    ws.cell(row=row_num, column=1, value=line.strip())
                    row_num += 1

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

# PDF 转图片
def pdf_to_images(pdf_file):

    pdf_bytes = pdf_file.read()

    pdf = fitz.open(
        stream=pdf_bytes,
        filetype="pdf"
    )

    zip_buffer = io.BytesIO()

    with ZipFile(zip_buffer, "a") as zip_file:

        for page_num in range(len(pdf)):

            page = pdf.load_page(page_num)

            pix = page.get_pixmap(
                matrix=fitz.Matrix(2, 2)
            )

            img_bytes = pix.tobytes("png")

            zip_file.writestr(
                f"page_{page_num+1}.png",
                img_bytes
            )

    zip_buffer.seek(0)

    return zip_buffer

if uploaded_files:
    if st.button("🚀 开始转换"):
        zip_buffer = io.BytesIO()
        if len(uploaded_files) > 1 or convert_option=="图片 (.png 每页)":
            zip_file = ZipFile(zip_buffer, "a")
        for file in uploaded_files:
            if convert_option == "Word (.docx)":
                result = pdf_to_word(file)
                st.download_button(
                    f"下载 {file.name[:-4]}.docx",
                    result,
                    file_name=file.name[:-4]+".docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                )
            elif convert_option == "Excel (.xlsx)":
                result = pdf_to_excel(file)
                st.download_button(
                    f"下载 {file.name[:-4]}.xlsx",
                    result,
                    file_name=file.name[:-4]+".xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )
            elif convert_option == "图片 (.png 每页)":
                result = pdf_to_images(file)
                zip_file.writestr(f"{file.name[:-4]}.zip", result.getvalue())
        if convert_option == "图片 (.png 每页)" and len(uploaded_files) > 0:
            zip_file.close()
            st.download_button(
                "下载所有图片 ZIP",
                zip_buffer,
                file_name="pdf_images.zip",
                mime="application/zip"
            )